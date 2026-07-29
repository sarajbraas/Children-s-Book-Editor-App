"""Manuscript ingestion: extract narration text from a raw .docx or .txt file.

Strips front/back matter (title, byline, email, wordcount, pitch paragraphs,
illustrator/publisher credits, page numbers, Project Gutenberg boilerplate)
and separates author illustration notes from the narration stream. Does not
chunk, tag, or otherwise interpret the text -- that happens in later phases.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import docx

GUTENBERG_START_RE = re.compile(r"\*\*\*\s*start of .*?project gutenberg.*?\*\*\*", re.IGNORECASE)
# Older Gutenberg transcriptions include a plain "End of Project Gutenberg's
# <title>, by <author>" sign-off line before the "*** END OF ... ***" marker.
GUTENBERG_END_RE = re.compile(
    r"(\*\*\*\s*end of .*?project gutenberg.*?\*\*\*|^end of .*project gutenberg.*$)",
    re.IGNORECASE | re.MULTILINE,
)

BLANK_LINE_SPLIT_RE = re.compile(r"\n\s*\n")

BARE_ILLUSTRATION_RE = re.compile(r"\[illustration:?\]", re.IGNORECASE)
ILLO_NOTE_RE = re.compile(r"\(\s*illo\s*:\s*(.*?)\)", re.IGNORECASE)
BARE_NUMBER_RE = re.compile(r"^\d+$")
PRODUCTION_NOTE_RE = re.compile(r"for reference only", re.IGNORECASE)
ATTRIBUTION_RE = re.compile(r"^(?:by|illustrated\s+by|illustrator)\s*:?\s*(.*)$", re.IGNORECASE)
BASED_ON_RE = re.compile(r"^based on\b", re.IGNORECASE)
EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+\.[\w.-]+$")
WORDCOUNT_RE = re.compile(r"^word\s*count\s*:", re.IGNORECASE)
PITCH_RE = re.compile(r"^(pitch|summary|synopsis)\s*:", re.IGNORECASE)
ILLUSTRATIONS_HEADER_RE = re.compile(r"^illustrations?$", re.IGNORECASE)


@dataclass
class IllustrationNote:
    """An author-provided illustrator/staging note pulled out of the narration."""

    position_after_line: int
    note_text: str


@dataclass
class IngestResult:
    narration_lines: list[str] = field(default_factory=list)
    illustration_notes: list[IllustrationNote] = field(default_factory=list)


def ingest(path) -> IngestResult:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        paragraphs = _paragraphs_from_docx(path)
    elif suffix == ".txt":
        paragraphs = _paragraphs_from_txt(path)
    else:
        raise ValueError(f"Unsupported manuscript format: {path.suffix}")

    narration_lines, illustration_notes = _extract_narration(paragraphs)
    return IngestResult(narration_lines=narration_lines, illustration_notes=illustration_notes)


def _paragraphs_from_docx(path: Path) -> list[str]:
    document = docx.Document(str(path))
    return [p.text for p in document.paragraphs]


def _paragraphs_from_txt(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    text = _strip_gutenberg_boilerplate(text)
    return BLANK_LINE_SPLIT_RE.split(text)


def _strip_gutenberg_boilerplate(text: str) -> str:
    start_match = GUTENBERG_START_RE.search(text)
    if start_match:
        text = text[start_match.end():]
    end_match = GUTENBERG_END_RE.search(text)
    if end_match:
        text = text[:end_match.start()]
    return text


PUBLISHER_SUFFIXES = {"co.", "co", "inc.", "inc", "ltd.", "ltd", "corp.", "corp"}


def _is_all_caps_line(text: str) -> bool:
    words = text.split()
    significant = [w for w in words if w.strip(".,").lower() not in PUBLISHER_SUFFIXES]
    letters = [c for c in " ".join(significant) if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


def _extract_illo_notes(text: str) -> tuple[str, list[str]]:
    notes = [m.group(1).strip() for m in ILLO_NOTE_RE.finditer(text)]
    cleaned = ILLO_NOTE_RE.sub("", text)
    # Bare "[Illustration]" placeholders (no descriptive text) carry no data.
    cleaned = BARE_ILLUSTRATION_RE.sub("", cleaned)
    cleaned = re.sub(r" {2,}", " ", cleaned).strip()
    return cleaned, notes


def _extract_narration(paragraphs: list[str]) -> tuple[list[str], list[IllustrationNote]]:
    narration_lines: list[str] = []
    illustration_notes: list[IllustrationNote] = []
    skip_next = False
    title_consumed = False

    for raw in paragraphs:
        if skip_next:
            skip_next = False
            continue

        text = raw.strip()
        if not text:
            continue

        # The manuscript's very first line of content is always its title.
        if not title_consumed:
            title_consumed = True
            continue

        # Title pages and publisher/imprint credits are conventionally set in
        # all caps; real read-aloud narration in these manuscripts never is.
        if _is_all_caps_line(text):
            continue
        if BARE_NUMBER_RE.match(text):
            continue
        if PRODUCTION_NOTE_RE.search(text):
            continue
        if WORDCOUNT_RE.match(text) or PITCH_RE.match(text) or BASED_ON_RE.match(text):
            continue
        if EMAIL_RE.match(text):
            continue
        if ILLUSTRATIONS_HEADER_RE.match(text):
            continue

        attribution = ATTRIBUTION_RE.match(text)
        if attribution:
            # A bare "By" (or "Illustrated by") with nothing else on the line
            # conventionally has the credited name on the following line.
            if not attribution.group(1).strip():
                skip_next = True
            continue

        cleaned, notes = _extract_illo_notes(raw)
        if cleaned.strip():
            narration_lines.append(cleaned)
        for note_text in notes:
            illustration_notes.append(
                IllustrationNote(position_after_line=len(narration_lines) - 1, note_text=note_text)
            )

    return narration_lines, illustration_notes
